"""يقرأ attendance_sample.xlsx ويحسب تأخير كل موظف.

قاعدة التأخير: أي حضور بعد الساعة 9:15 صباحاً يُعد تأخيراً.
الحضور في 9:15 تماماً أو قبلها لا يُحسب.

التشغيل من مجلد attendance:
    python late_report.py

الناتج: late_summary.xlsx مرتباً من الأكثر تأخيراً إلى الأقل.
"""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SOURCE_PATH = HERE / "attendance_sample.xlsx"
OUTPUT_PATH = HERE / "late_summary.xlsx"

# الحضور بعد هذه الساعة فقط هو التأخير.
THRESHOLD = time(9, 15)

SOURCE_HEADERS = [
    "كود الموظف",
    "الاسم",
    "القسم",
    "التاريخ",
    "وقت الحضور",
    "وقت الانصراف",
]

SUMMARY_HEADERS = [
    "الترتيب",
    "كود الموظف",
    "الاسم",
    "القسم",
    "عدد أيام التأخير",
    "إجمالي دقائق التأخير",
]


def as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"تاريخ غير مفهوم: {value!r}")


def as_time(value: object) -> time:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(text, pattern).time()
            except ValueError:
                continue
    raise ValueError(f"وقت غير مفهوم: {value!r}")


def late_minutes(check_in: time) -> int:
    """دقائق التأخير بعد 9:15. الصفر يعني الموظف في الموعد.

    أي جزء من الدقيقة بعد 9:15 يُقرّب إلى دقيقة كاملة،
    لأن الحضور «بعد» الحد حتى بثوانٍ ما زال تأخيراً.
    """
    arrived = check_in.hour * 3600 + check_in.minute * 60 + check_in.second
    limit = THRESHOLD.hour * 3600 + THRESHOLD.minute * 60 + THRESHOLD.second
    if arrived <= limit:
        return 0
    return (arrived - limit + 59) // 60


def load_attendance(path: Path) -> list[dict]:
    if not path.exists():
        raise FileNotFoundError(
            f"الملف غير موجود: {path}\nشغّل generate_attendance.py أولاً."
        )

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    if headers[: len(SOURCE_HEADERS)] != SOURCE_HEADERS:
        raise ValueError(
            "عناوين الملف غير مطابقة. المتوقع في الصف الأول: "
            + "، ".join(SOURCE_HEADERS)
        )

    rows: list[dict] = []
    seen: set[tuple[str, date]] = set()
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        if excel_row[0] in (None, ""):
            continue
        code = str(excel_row[0]).strip()
        name = str(excel_row[1]).strip()
        department = str(excel_row[2]).strip()
        day = as_date(excel_row[3])
        check_in = as_time(excel_row[4])
        # وقت الانصراف موجود للمراجعة، والتأخير يُحسب من وقت الحضور فقط.
        as_time(excel_row[5])
        key = (code, day)
        if key in seen:
            raise ValueError(f"صف مكرر للموظف {code} في تاريخ {day}")
        seen.add(key)
        rows.append(
            {
                "code": code,
                "name": name,
                "department": department,
                "day": day,
                "check_in": check_in,
            }
        )

    if not rows:
        raise ValueError("ملف الحضور لا يحتوي على صفوف بيانات.")
    return rows


def summarize(rows: list[dict]) -> list[dict]:
    """يجمع أيام التأخير ودقائقها لكل موظف، ثم يرتبها من الأكثر إلى الأقل."""
    grouped: dict[str, dict] = {}
    for row in rows:
        item = grouped.setdefault(
            row["code"],
            {
                "code": row["code"],
                "name": row["name"],
                "department": row["department"],
                "late_days": 0,
                "late_minutes": 0,
            },
        )
        minutes = late_minutes(row["check_in"])
        if minutes > 0:
            item["late_days"] += 1
            item["late_minutes"] += minutes

    return sorted(
        grouped.values(),
        key=lambda item: (-item["late_minutes"], -item["late_days"], item["code"]),
    )


def _styles():
    thin = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )
    return {
        "thin": thin,
        "title_font": Font(name="Calibri", bold=True, color="1F4E79", size=18),
        "note_font": Font(name="Calibri", color="526072", size=11),
        "header_font": Font(name="Calibri", bold=True, color="FFFFFF", size=12),
        "header_fill": PatternFill("solid", fgColor="1F4E79"),
        "cell_font": Font(name="Calibri", size=12, color="1F2933"),
        "late_font": Font(name="Calibri", bold=True, size=12, color="9C2F2F"),
        "ontime_font": Font(name="Calibri", bold=True, size=12, color="1E7A46"),
        "total_font": Font(name="Calibri", bold=True, size=12, color="1F4E79"),
        "title_fill": PatternFill("solid", fgColor="E8EEF7"),
        "note_fill": PatternFill("solid", fgColor="F7F9FC"),
        "zebra": PatternFill("solid", fgColor="F7F9FC"),
        "ontime_fill": PatternFill("solid", fgColor="F3FBF6"),
        "total_fill": PatternFill("solid", fgColor="E8EEF7"),
        "center": Alignment(horizontal="center", vertical="center"),
        "right": Alignment(horizontal="right", vertical="center"),
    }


def write_summary(ranked: list[dict], source_rows: list[dict], path: Path) -> None:
    days = sorted({row["day"] for row in source_rows})
    first_day = days[0].strftime("%d/%m/%Y")
    last_day = days[-1].strftime("%d/%m/%Y")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ملخص التأخير"
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    styles = _styles()

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "تقرير التأخير الأسبوعي"
    for column in range(1, 7):
        cell = sheet.cell(1, column)
        cell.font = styles["title_font"]
        cell.fill = styles["title_fill"]
        cell.alignment = styles["right"]

    sheet.merge_cells("A2:F2")
    sheet["A2"] = (
        f"يُحتسب التأخير عند الحضور بعد 9:15 صباحاً  |  "
        f"الفترة: {first_day} إلى {last_day}  |  "
        f"الترتيب حسب إجمالي الدقائق ثم عدد الأيام"
    )
    for column in range(1, 7):
        cell = sheet.cell(2, column)
        cell.font = styles["note_font"]
        cell.fill = styles["note_fill"]
        cell.alignment = styles["right"]

    header_row = 4
    for column, header in enumerate(SUMMARY_HEADERS, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["center"]
        cell.border = styles["thin"]

    for offset, item in enumerate(ranked, start=1):
        row_index = header_row + offset
        values = [
            offset,
            item["code"],
            item["name"],
            item["department"],
            item["late_days"],
            item["late_minutes"],
        ]
        on_time = item["late_days"] == 0
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = styles["cell_font"]
            cell.alignment = styles["center"]
            cell.border = styles["thin"]
            if on_time:
                cell.fill = styles["ontime_fill"]
            elif offset % 2 == 0:
                cell.fill = styles["zebra"]
        minutes_cell = sheet.cell(row_index, 6)
        days_cell = sheet.cell(row_index, 5)
        minutes_cell.number_format = "#,##0"
        days_cell.number_format = "0"
        if on_time:
            minutes_cell.font = styles["ontime_font"]
            days_cell.font = styles["ontime_font"]
        else:
            minutes_cell.font = styles["late_font"]
            days_cell.font = styles["late_font"]
        sheet.row_dimensions[row_index].height = 22

    total_row = header_row + len(ranked) + 1
    totals = [
        "",
        "",
        "",
        "الإجمالي",
        sum(item["late_days"] for item in ranked),
        sum(item["late_minutes"] for item in ranked),
    ]
    for column, value in enumerate(totals, start=1):
        cell = sheet.cell(total_row, column, value)
        cell.font = styles["total_font"]
        cell.fill = styles["total_fill"]
        cell.alignment = styles["center"]
        cell.border = styles["thin"]
    sheet.cell(total_row, 6).number_format = "#,##0"

    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[header_row].height = 24
    sheet.row_dimensions[total_row].height = 24
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A{header_row}:F{header_row + len(ranked)}"
    sheet.print_title_rows = "1:4"

    widths = [12, 16, 28, 22, 22, 26]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.right.text = "حضور بعد 9:15 = تأخير"
    sheet.oddFooter.left.text = "صفحة &P من &N"

    workbook.properties.title = "ملخص التأخير"
    workbook.properties.creator = "late_report.py"
    workbook.save(path)


def main() -> None:
    rows = load_attendance(SOURCE_PATH)
    ranked = summarize(rows)
    write_summary(ranked, rows, OUTPUT_PATH)

    late_people = sum(1 for item in ranked if item["late_days"] > 0)
    print(f"تم قراءة: {SOURCE_PATH}")
    print(f"عدد صفوف الحضور: {len(rows)}")
    print(f"عدد الموظفين: {len(ranked)}")
    print(f"موظفون لديهم تأخير: {late_people}")
    print(f"تم إنشاء الملف: {OUTPUT_PATH}")
    print("الأكثر تأخيراً:")
    for item in ranked[:5]:
        print(
            f"  {item['code']} {item['name']}: "
            f"{item['late_days']} يوم، {item['late_minutes']} دقيقة"
        )


if __name__ == "__main__":
    main()
