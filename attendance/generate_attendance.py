"""توليد ملف حضور وهمي لأسبوع عمل من الأحد إلى الخميس.

التشغيل من مجلد attendance:
    python generate_attendance.py

الناتج: attendance_sample.xlsx في نفس المجلد.
البيانات كلها وهمية، والأرقام مثبتة ببذرة عشوائية حتى تتكرر نفس النتيجة.
"""

from __future__ import annotations

import random
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# بذرة ثابتة: نفس الملف يتولد في كل تشغيل.
RANDOM_SEED = 42
# الأحد 11 أكتوبر 2026، ثم الإثنين حتى الخميس.
WEEK_START = date(2026, 10, 11)
OUTPUT_PATH = Path(__file__).resolve().parent / "attendance_sample.xlsx"

HEADERS = [
    "كود الموظف",
    "الاسم",
    "القسم",
    "التاريخ",
    "وقت الحضور",
    "وقت الانصراف",
]

# (الكود، الاسم الوهمي، القسم، عدد أيام التأخير من أصل 5)
# بعض الموظفين لا يتأخرون، وبعضهم يتأخر في يوم أو أكثر.
EMPLOYEES = [
    ("EMP001", "أحمد محمود حسن", "الموارد البشرية", 0),
    ("EMP002", "فاطمة علي إبراهيم", "الموارد البشرية", 0),
    ("EMP003", "محمد سعيد عبد الله", "المالية", 0),
    ("EMP004", "نورا حسن محمد", "المالية", 1),
    ("EMP005", "يوسف كمال أحمد", "تقنية المعلومات", 1),
    ("EMP006", "مريم عبد الرحمن سالم", "تقنية المعلومات", 1),
    ("EMP007", "عمر فاروق السيد", "المبيعات", 2),
    ("EMP008", "سلمى جمال الدين", "المبيعات", 2),
    ("EMP009", "خالد إبراهيم ناصر", "العمليات", 2),
    ("EMP010", "هدى مصطفى علي", "العمليات", 0),
    ("EMP011", "طارق عبد العزيز", "الموارد البشرية", 3),
    ("EMP012", "ليلى سمير حسن", "المالية", 3),
    ("EMP013", "حسن محمد فؤاد", "تقنية المعلومات", 1),
    ("EMP014", "رنا عادل كريم", "المبيعات", 4),
    ("EMP015", "إبراهيم سامي نور", "العمليات", 2),
    ("EMP016", "دينا وليد عبد الله", "الموارد البشرية", 0),
    ("EMP017", "مصطفى حسام الدين", "المالية", 3),
    ("EMP018", "آمال فتحي محمود", "تقنية المعلومات", 5),
    ("EMP019", "كريم ياسر إبراهيم", "المبيعات", 4),
    ("EMP020", "شيماء عادل حسن", "العمليات", 1),
]


def work_week(start: date) -> list[date]:
    """خمسة أيام عمل متتالية تبدأ من الأحد."""
    return [start + timedelta(days=offset) for offset in range(5)]


def minutes_to_time(total_minutes: int) -> time:
    return time(total_minutes // 60, total_minutes % 60)


def build_rows(rng: random.Random) -> list[dict]:
    """يبني صفاً لكل موظف في كل يوم عمل."""
    days = work_week(WEEK_START)
    rows: list[dict] = []

    for code, name, department, late_day_count in EMPLOYEES:
        late_indexes = set(rng.sample(range(len(days)), late_day_count))
        for index, day in enumerate(days):
            if index in late_indexes:
                # بعد 9:15: من 9:16 حتى 10:30.
                check_in = minutes_to_time(rng.randint(9 * 60 + 16, 10 * 60 + 30))
            else:
                # في الموعد: من 8:20 حتى 9:15 شاملة.
                check_in = minutes_to_time(rng.randint(8 * 60 + 20, 9 * 60 + 15))
            check_out = minutes_to_time(rng.randint(16 * 60, 17 * 60 + 15))
            rows.append(
                {
                    "كود الموظف": code,
                    "الاسم": name,
                    "القسم": department,
                    "التاريخ": day,
                    "وقت الحضور": check_in,
                    "وقت الانصراف": check_out,
                }
            )

    # مثالان واضحان لحد التأخير: 9:15 في الموعد، و 9:16 متأخر بدقيقة.
    _set_check_in(rows, "EMP001", days[0], time(9, 15))
    _set_check_in(rows, "EMP004", late_day_for(rows, "EMP004", days), time(9, 16))
    return rows


def late_day_for(rows: list[dict], code: str, days: list[date]) -> date:
    """أول يوم حضوره بعد 9:15 لهذا الموظف، وهو اليوم الذي سنثبّت عليه 9:16."""
    for row in rows:
        if row["كود الموظف"] != code:
            continue
        check_in: time = row["وقت الحضور"]
        if (check_in.hour, check_in.minute) > (9, 15):
            return row["التاريخ"]
    return days[0]


def _set_check_in(rows: list[dict], code: str, day: date, check_in: time) -> None:
    for row in rows:
        if row["كود الموظف"] == code and row["التاريخ"] == day:
            row["وقت الحضور"] = check_in
            return
    raise KeyError(f"لا يوجد صف للموظف {code} في {day}")


def write_workbook(rows: list[dict], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الحضور"
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    cell_font = Font(name="Calibri", size=12)
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )
    zebra = PatternFill("solid", fgColor="F7F9FC")

    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row[header] for header in HEADERS])

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    for row_index in range(2, sheet.max_row + 1):
        for column_index in range(1, len(HEADERS) + 1):
            cell = sheet.cell(row_index, column_index)
            cell.font = cell_font
            cell.alignment = center
            cell.border = thin
            if row_index % 2 == 0:
                cell.fill = zebra
        sheet.cell(row_index, 4).number_format = "DD/MM/YYYY"
        sheet.cell(row_index, 5).number_format = "HH:MM"
        sheet.cell(row_index, 6).number_format = "HH:MM"
        sheet.row_dimensions[row_index].height = 20

    sheet.row_dimensions[1].height = 24
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    sheet.freeze_panes = "A2"

    widths = [16, 28, 22, 16, 16, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.center.text = "سجل حضور وهمي — أسبوع عمل"
    sheet.print_title_rows = "1:1"

    workbook.properties.title = "سجل حضور وهمي"
    workbook.properties.creator = "generate_attendance.py"
    workbook.save(path)


def main() -> None:
    rng = random.Random(RANDOM_SEED)
    rows = build_rows(rng)
    write_workbook(rows, OUTPUT_PATH)
    days = work_week(WEEK_START)
    print(f"تم إنشاء الملف: {OUTPUT_PATH}")
    print(f"عدد الموظفين: {len(EMPLOYEES)}")
    print(f"الأسبوع: {days[0].strftime('%d/%m/%Y')} إلى {days[-1].strftime('%d/%m/%Y')} (الأحد–الخميس)")
    print(f"عدد الصفوف: {len(rows)}")


if __name__ == "__main__":
    main()
